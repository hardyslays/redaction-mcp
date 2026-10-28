import base64
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from io import BytesIO
from pathlib import Path
import sqlite3
import threading
import time

import pytest
import pymupdf as fitz
from docx import Document as DocxDocument

from redaction_mcp.core.models.jobs import JobSubmit, UploadCreate
from redaction_mcp.core.services.mime_detection import DOCX_MIME_TYPE
from redaction_mcp.core.services.jobs import JobConflictError, JobManager
from redaction_mcp.core.services.limits import CHUNK_SIZE, DIRECT_LIMIT, OFFICE_LIMIT


def _completed_job(manager: JobManager):
    data = b"name,email\nJane Doe,jane@example.com\n"
    upload = manager.create_upload(UploadCreate(filename="people.csv", size_bytes=len(data)))
    manager.upload_chunk(upload.upload_id, 0, data)
    job = manager.submit_job(
        JobSubmit(upload_id=upload.upload_id, operation="redact", targets=[{"type": "text", "values": ["Jane Doe"]}])
    )
    deadline = time.monotonic() + 5
    while job.status not in {"succeeded", "failed"} and time.monotonic() < deadline:
        time.sleep(0.01)
        job = manager.get_job(job.job_id)
    return job


def test_job_chunks_process_persist_and_download(tmp_path: Path) -> None:
    manager = JobManager(tmp_path)
    job = _completed_job(manager)

    assert job.status == "succeeded"
    chunk = manager.result_chunk(job.job_id, 0, CHUNK_SIZE)
    assert base64.b64decode(chunk.base64data) == b"name,email\n**** ***,jane@example.com\n"
    assert chunk.eof is True
    assert JobManager(tmp_path).get_job(job.job_id).status == "succeeded"
    manager.stop()


def test_job_rejects_out_of_order_chunks_incomplete_uploads_and_large_office(tmp_path: Path) -> None:
    manager = JobManager(tmp_path)
    upload = manager.create_upload(UploadCreate(filename="people.csv", size_bytes=2))
    with pytest.raises(JobConflictError, match="Expected chunk index 0"):
        manager.upload_chunk(upload.upload_id, 1, b"x")
    manager.upload_chunk(upload.upload_id, 0, b"x")
    with pytest.raises(JobConflictError, match="incomplete"):
        manager.submit_job(JobSubmit(upload_id=upload.upload_id, operation="redact", targets=[{"type": "text", "values": ["x"]}]))
    with pytest.raises(ValueError, match="limited"):
        manager.create_upload(UploadCreate(filename="large.xlsx", size_bytes=OFFICE_LIMIT + 1))
    with pytest.raises(ValueError, match="limited"):
        manager.create_upload(UploadCreate(filename="large.txt", size_bytes=DIRECT_LIMIT + 1))
    manager.stop()


def test_concurrent_uploads_cannot_append_same_chunk_twice(tmp_path: Path) -> None:
    manager = JobManager(tmp_path)
    upload = manager.create_upload(UploadCreate(filename="people.csv", size_bytes=1))
    barrier = threading.Barrier(2)

    def send() -> str:
        barrier.wait()
        try:
            manager.upload_chunk(upload.upload_id, 0, b"x")
            return "accepted"
        except JobConflictError:
            return "conflict"

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: send(), range(2)))

    assert sorted(results) == ["accepted", "conflict"]
    assert (manager.upload_path / f"{upload.upload_id}.bin").read_bytes() == b"x"
    manager.stop()


def test_office_job_uses_engine_without_direct_request_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import redaction_mcp.core.services.jobs as jobs

    source = DocxDocument()
    source.add_paragraph("secret visible")
    buffer = BytesIO()
    source.save(buffer)
    data = buffer.getvalue()
    manager = JobManager(tmp_path)
    upload = manager.create_upload(UploadCreate(filename="source.docx", size_bytes=len(data)))
    manager.upload_chunk(upload.upload_id, 0, data)

    def direct_request(*args: object, **kwargs: object) -> None:
        raise AssertionError("Office jobs must bypass the direct request limit")

    monkeypatch.setattr(jobs, "redact_document", direct_request)
    job = manager.submit_job(JobSubmit(
        upload_id=upload.upload_id, operation="redact",
        targets=[{"type": "text", "values": ["secret"]}],
    ))
    deadline = time.monotonic() + 5
    while job.status not in {"succeeded", "failed"} and time.monotonic() < deadline:
        time.sleep(0.01)
        job = manager.get_job(job.job_id)

    assert job.status == "succeeded"
    output = DocxDocument(BytesIO(manager.result_bytes(job.job_id, 0, CHUNK_SIZE)[0]))
    assert output.paragraphs[0].text == "****** visible"
    manager.stop()


def test_office_mime_requires_archive_validation(tmp_path: Path) -> None:
    manager = JobManager(tmp_path)
    upload = manager.create_upload(UploadCreate(filename="source.bin", mime_type=DOCX_MIME_TYPE, size_bytes=3))
    manager.upload_chunk(upload.upload_id, 0, b"bad")

    with pytest.raises(ValueError, match="valid ZIP"):
        manager.submit_job(JobSubmit(
            upload_id=upload.upload_id, operation="redact",
            targets=[{"type": "text", "values": ["secret"]}],
        ))
    manager.stop()


def test_restart_marks_running_failed_and_cleanup_removes_terminal_files(tmp_path: Path) -> None:
    manager = JobManager(tmp_path)
    job = _completed_job(manager)
    with closing(sqlite3.connect(manager.database_path)) as connection:
        with connection:
            connection.execute("UPDATE jobs SET status='running' WHERE id=?", (job.job_id,))
    restarted = JobManager(tmp_path, retention_seconds=0)
    assert restarted.get_job(job.job_id).status == "failed"
    assert restarted.cleanup() == 1
    assert not (restarted.result_path / f"{job.job_id}.bin").exists()
    manager.stop()
    restarted.stop()


def test_pdf_job_uses_file_backed_processing(tmp_path: Path) -> None:
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "secret visible")
    data = pdf.tobytes()
    pdf.close()
    manager = JobManager(tmp_path)
    upload = manager.create_upload(UploadCreate(filename="source.pdf", size_bytes=len(data)))
    manager.upload_chunk(upload.upload_id, 0, data)
    job = manager.submit_job(
        JobSubmit(upload_id=upload.upload_id, operation="redact", targets=[{"type": "text", "values": ["secret"]}])
    )
    deadline = time.monotonic() + 5
    while job.status not in {"succeeded", "failed"} and time.monotonic() < deadline:
        time.sleep(0.01)
        job = manager.get_job(job.job_id)

    assert job.status == "succeeded"
    output = fitz.open(stream=manager.result_bytes(job.job_id, 0, CHUNK_SIZE)[0], filetype="pdf")
    assert "secret" not in output[0].get_text()
    assert "visible" in output[0].get_text()
    output.close()
    manager.stop()
