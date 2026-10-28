import base64
import csv
from copy import copy
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook

from redaction_mcp.core.models.document import Document
from redaction_mcp.core.models.errors import RedactionError
from redaction_mcp.core.models.redaction import BoundingBox, BoundingBoxTarget, RegexTarget, TextTarget
from redaction_mcp.core.models.replacement import TextReplacementTarget
from redaction_mcp.core.services.redaction_service import redact_document
from redaction_mcp.core.services.replacement_service import replace_document


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "People"
    sheet["A1"] = "Jane Doe and public"
    font = copy(sheet["A1"].font)
    font.bold = True
    sheet["A1"].font = font
    sheet["B1"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_redacts_partial_cells_and_preserves_formula_style_and_sheet() -> None:
    result = redact_document(
        Document(base64data=base64.b64encode(_xlsx()).decode(), filename="people.xlsx"),
        [TextTarget(type="text", values=["Jane Doe"])],
    )

    assert isinstance(result, Document)
    workbook = load_workbook(BytesIO(base64.b64decode(result.base64data)), data_only=False)
    assert workbook.sheetnames == ["People"]
    assert workbook["People"]["A1"].value == "**** *** and public"
    assert workbook["People"]["A1"].font.bold is True
    assert workbook["People"]["B1"].value == "=1+1"


def test_xlsx_replacement_and_unsupported_geometry() -> None:
    replaced = replace_document(
        Document(base64data=base64.b64encode(_xlsx()).decode(), filename="people.xlsx"),
        [TextReplacementTarget(type="text", values=["Jane Doe"], replacement_type="STATIC", static_text="SAFE")],
    )
    assert isinstance(replaced, Document)
    workbook = load_workbook(BytesIO(base64.b64decode(replaced.base64data)))
    assert workbook.active["A1"].value == "[SAFE] and public"

    rejected = redact_document(
        Document(base64data=base64.b64encode(_xlsx()).decode(), filename="people.xlsx"),
        [BoundingBoxTarget(type="bounding_box", values=[BoundingBox(page=0, x=0, y=0, width=1, height=1)])],
    )
    assert isinstance(rejected, RedactionError)
    assert rejected.message == "Spreadsheet redaction supports only text and regex targets"


def test_csv_preserves_bom_delimiter_newlines_and_supports_regex() -> None:
    data = "\ufeffname;id\r\nJane Doe;AB123456\r\n".encode()
    result = redact_document(
        Document(base64data=base64.b64encode(data).decode(), filename="people.csv"),
        [RegexTarget(type="regex", patterns=[r"[A-Z]{2}\d{6}"], ignore_case=False)],
    )

    assert isinstance(result, Document)
    output = base64.b64decode(result.base64data)
    assert output.startswith(b"\xef\xbb\xbf")
    assert b";" in output and b"\r\n" in output
    assert b"AB123456" not in output


def test_csv_multiline_cell_redaction_keeps_quoted_row_structure() -> None:
    data = b'name,address\nJane,"123 Main\nStreet"\n'
    result = redact_document(
        Document(base64data=base64.b64encode(data).decode(), filename="people.csv"),
        [TextTarget(type="text", values=["Main Street"])],
    )

    assert isinstance(result, Document)
    rows = list(csv.reader(StringIO(base64.b64decode(result.base64data).decode())))
    assert rows == [["name", "address"], ["Jane", "123 ****\n******"]]


def test_csv_semantic_replacement_is_stable_for_repeated_values() -> None:
    data = b"name\nJane Doe\nJane Doe\n"
    result = replace_document(
        Document(base64data=base64.b64encode(data).decode(), filename="people.csv"),
        [TextReplacementTarget(type="text", values=["Jane Doe"], replacement_type="SEMANTIC", entity_type="NAME")],
    )

    assert isinstance(result, Document)
    lines = base64.b64decode(result.base64data).decode().splitlines()
    assert lines[1] == lines[2]
    assert lines[1] != "Jane Doe"
