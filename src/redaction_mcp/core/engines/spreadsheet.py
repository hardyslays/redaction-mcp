"""XLSX and CSV redaction and replacement engines."""

from __future__ import annotations

import base64
import csv
from io import BytesIO, StringIO
from pathlib import Path
import re

from openpyxl import load_workbook

from redaction_mcp.core.models.document import Document
from redaction_mcp.core.models.redaction import (
    BoundingBoxTarget,
    PageTarget,
    RedactionOptions,
    RedactionTarget,
    RegexTarget,
    TextTarget,
)
from redaction_mcp.core.models.replacement import ReplacementTarget
from redaction_mcp.core.services.replacement_text import ReplacementContext
from redaction_mcp.core.services.text_matching import replace_text_matches

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME_TYPE = "text/csv"


def _mask_text(value: str, options: RedactionOptions) -> str:
    if options.redaction_type == "mask":
        return "[REDACT]"
    return "".join(char if char.isspace() else "*" for char in value)


def _redact_value(value: str, target: RedactionTarget, options: RedactionOptions) -> str:
    replacement = lambda match: _mask_text(match.group(), options)
    if isinstance(target, TextTarget):
        if target.pages is not None:
            raise ValueError("Spreadsheet redaction does not support page-restricted targets")
        return replace_text_matches(
            value,
            target.values,
            replacement,
            ignore_case=target.ignore_case,
            partial_match=target.partial_match,
        )[0]
    if isinstance(target, RegexTarget):
        if target.pages is not None:
            raise ValueError("Spreadsheet redaction does not support page-restricted targets")
        flags = re.IGNORECASE if target.ignore_case else 0
        if not target.allow_unicode:
            flags |= re.ASCII
        result = value
        for pattern in target.patterns:
            result = re.sub(
                pattern,
                replacement,
                result,
                count=1 if target.only_first_match else 0,
                flags=flags,
            )
            if target.only_first_match and result != value:
                break
        return result
    if isinstance(target, (BoundingBoxTarget, PageTarget)):
        raise ValueError("Spreadsheet redaction supports only text and regex targets")
    return value


def _replace_value(value: str, target: ReplacementTarget, context: ReplacementContext) -> str:
    if target.pages is not None:
        raise ValueError("Spreadsheet replacement does not support page-restricted targets")
    return replace_text_matches(
        value,
        target.values,
        lambda match: context.replacement_text(match.group(), target),
        ignore_case=target.ignore_case,
        partial_match=target.partial_match,
    )[0]


def _xlsx_cells(workbook: object):
    for worksheet in workbook.worksheets:  # type: ignore[attr-defined]
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "s" and isinstance(cell.value, str):
                    yield cell


def redact_xlsx_document(
    document: Document, targets: list[RedactionTarget], options: RedactionOptions, *, data: bytes | None = None
) -> Document:
    workbook = load_workbook(BytesIO(data if data is not None else document.decoded_bytes()), data_only=False)
    for cell in _xlsx_cells(workbook):
        for target in targets:
            cell.value = _redact_value(cell.value, target, options)
    output = BytesIO()
    workbook.save(output)
    filename = Path(document.filename or "document.xlsx")
    return Document(
        base64data=base64.b64encode(output.getvalue()).decode("ascii"),
        mime_type=XLSX_MIME_TYPE,
        filename=f"{filename.stem}-redacted.xlsx",
    )


def replace_xlsx_document(
    document: Document, targets: list[ReplacementTarget], *, data: bytes | None = None
) -> Document:
    workbook = load_workbook(BytesIO(data if data is not None else document.decoded_bytes()), data_only=False)
    context = ReplacementContext()
    for cell in _xlsx_cells(workbook):
        for target in targets:
            cell.value = _replace_value(cell.value, target, context)
    output = BytesIO()
    workbook.save(output)
    filename = Path(document.filename or "document.xlsx")
    return Document(
        base64data=base64.b64encode(output.getvalue()).decode("ascii"),
        mime_type=XLSX_MIME_TYPE,
        filename=f"{filename.stem}-replaced.xlsx",
    )


def _csv_format(data: bytes) -> tuple[str, csv.Dialect, str, bool]:
    bom = data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    newline = "\r\n" if "\r\n" in text else "\n"
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return text, dialect, newline, bom


def _transform_csv(data: bytes, transform: object) -> bytes:
    text, dialect, newline, bom = _csv_format(data)
    output = StringIO(newline="")
    writer = csv.writer(output, dialect=dialect, lineterminator=newline)
    for row in csv.reader(StringIO(text, newline=""), dialect=dialect):
        writer.writerow([transform(value) for value in row])  # type: ignore[operator]
    encoded = output.getvalue().encode("utf-8")
    return (b"\xef\xbb\xbf" + encoded) if bom else encoded


def redact_csv_document(
    document: Document, targets: list[RedactionTarget], options: RedactionOptions, *, data: bytes | None = None
) -> Document:
    source = data if data is not None else document.decoded_bytes()

    def transform(value: str) -> str:
        for target in targets:
            value = _redact_value(value, target, options)
        return value

    output = _transform_csv(source, transform)
    filename = Path(document.filename or "document.csv")
    return Document(
        base64data=base64.b64encode(output).decode("ascii"), mime_type=CSV_MIME_TYPE, filename=f"{filename.stem}-redacted.csv"
    )


def replace_csv_document(
    document: Document, targets: list[ReplacementTarget], *, data: bytes | None = None
) -> Document:
    source = data if data is not None else document.decoded_bytes()
    context = ReplacementContext()

    def transform(value: str) -> str:
        for target in targets:
            value = _replace_value(value, target, context)
        return value

    output = _transform_csv(source, transform)
    filename = Path(document.filename or "document.csv")
    return Document(
        base64data=base64.b64encode(output).decode("ascii"), mime_type=CSV_MIME_TYPE, filename=f"{filename.stem}-replaced.csv"
    )
