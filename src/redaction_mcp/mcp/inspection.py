"""Bounded document-text extraction for MCP agents."""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from io import BytesIO, StringIO

import pymupdf as fitz
from docx import Document as DocxDocument
from pydantic import BaseModel, Field
from pptx import Presentation
from openpyxl import load_workbook
import csv

from redaction_mcp.core.engines.docx_redaction import _paragraphs
from redaction_mcp.core.engines.pptx_redaction import _text_frames
from redaction_mcp.core.models.document import Document
from redaction_mcp.core.services.limits import DIRECT_LIMIT
from redaction_mcp.core.services.mime_detection import (
    DOCX_MIME_TYPE,
    PDF_MIME_TYPE,
    PPTM_MIME_TYPE,
    PPTX_MIME_TYPE,
    TEXT_MIME_TYPE,
    CSV_MIME_TYPE,
    XLSX_MIME_TYPE,
    detect_mime_type,
)


class DocumentPage(BaseModel):
    """Text visible to an agent on one zero-based document page or slide."""

    index: int = Field(ge=0)
    text: str


class DocumentInspection(BaseModel):
    """Bounded text extraction that lets agents select redaction targets."""

    filename: str
    mime_type: str
    pages: list[DocumentPage]
    truncated: bool


def _joined(parts: Iterable[str], separator: str = "\n") -> Iterator[str]:
    first = True
    for part in parts:
        if not first:
            yield separator
        yield part
        first = False


def _bounded_pages(
    contents: Iterable[tuple[int, Iterable[str]]], max_characters: int
) -> tuple[list[DocumentPage], bool]:
    pages: list[DocumentPage] = []
    remaining = max_characters
    for index, pieces in contents:
        if remaining <= 0:
            return pages, True
        collected: list[str] = []
        for piece in pieces:
            if len(piece) > remaining:
                collected.append(piece[:remaining])
                pages.append(DocumentPage(index=index, text="".join(collected)))
                return pages, True
            collected.append(piece)
            remaining -= len(piece)
        pages.append(DocumentPage(index=index, text="".join(collected)))
    return pages, False


def inspect_document(document: Document, max_characters: int) -> DocumentInspection:
    """Extract visible text without returning more than ``max_characters``."""
    if max_characters < 1:
        raise ValueError("max_characters must be positive")
    if len(document.base64data) > ((DIRECT_LIMIT + 2) // 3) * 4:
        raise ValueError("Document inspection is limited to 16 MiB")
    data = document.decoded_bytes()
    if len(data) > DIRECT_LIMIT:
        raise ValueError("Document inspection is limited to 16 MiB")
    mime_type = document.mime_type or detect_mime_type(data, document.filename)
    if mime_type == PDF_MIME_TYPE:
        pdf = fitz.open(stream=data, filetype="pdf")
        try:
            pages, truncated = _bounded_pages(
                ((index, (page.get_text(),)) for index, page in enumerate(pdf)), max_characters
            )
        finally:
            pdf.close()
    elif mime_type == DOCX_MIME_TYPE:
        source = DocxDocument(BytesIO(data))
        pages, truncated = _bounded_pages(
            ((0, _joined(paragraph.text for paragraph in _paragraphs(source))),), max_characters
        )
    elif mime_type in (PPTX_MIME_TYPE, PPTM_MIME_TYPE):
        presentation = Presentation(BytesIO(data))
        pages, truncated = _bounded_pages(
            (
                (index, _joined(frame.text for frame in _text_frames(slide.shapes)))
                for index, slide in enumerate(presentation.slides)
            ), max_characters
        )
    elif mime_type == TEXT_MIME_TYPE:
        pages, truncated = _bounded_pages(((0, (data.decode("utf-8"),)),), max_characters)
    elif mime_type == CSV_MIME_TYPE:
        rows = csv.reader(StringIO(data.decode("utf-8-sig"), newline=""))
        pages, truncated = _bounded_pages(((0, _joined(" | ".join(row) for row in rows)),), max_characters)
    elif mime_type == XLSX_MIME_TYPE:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            pages, truncated = _bounded_pages(
                (
                    (
                        index,
                        _joined(
                            " | ".join(str(value) if value is not None else "" for value in row)
                            for row in worksheet.iter_rows(values_only=True)
                        ),
                    )
                    for index, worksheet in enumerate(workbook.worksheets)
                ),
                max_characters,
            )
        finally:
            workbook.close()
    else:
        raise ValueError("Only PDF, DOCX, PPTX, PPTM, TXT, CSV, and XLSX documents are currently supported")
    return DocumentInspection(
        filename=document.filename or "document",
        mime_type=mime_type or "application/octet-stream",
        pages=pages,
        truncated=truncated,
    )
